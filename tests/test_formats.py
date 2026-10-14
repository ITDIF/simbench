"""点表格式表（`project/formats.py`）：加格式只改一处，且三种格式语义一致。

这组用例同时是"接一种新格式要写什么"的说明书：一个 `Format` 就是
`(后缀, 对话框里那一行文字, 读, 写)` 四样，加上读/写的两条口径（读不懂抛 `ValueError`、
文件读不到抛 `OSError`）。
"""

from __future__ import annotations

import pathlib

import pytest

from simbench.csv_io_entry import load_point_table
from simbench.gui import dialogs
from simbench.project import csv_io, formats, json_io, xlsx_io
from simbench.project.model import Project, TcpChannel, Device, Point

ROOT = pathlib.Path(__file__).resolve().parents[1]
CSV = str(ROOT / "examples" / "covi-tunnel.csv")


def example() -> Project:
    project, _notes = formats.read_project(CSV)
    return project


def rows(project) -> list[list[str]]:
    return csv_io.rows_of(project)


# ---------- 格式表本身 ----------

def test_the_registry_lists_the_table_formats_before_the_project_one():
    assert [(fmt.suffix, fmt.kind) for fmt in formats.FORMATS] == [
        (".csv", "table"), (".xlsx", "table"), (".json", "project")]


def test_an_unknown_suffix_is_refused_with_the_supported_list():
    """后缀不认识就拒绝，别猜：猜错的表现是"把 JSON 内容当 CSV 解析"那种离题万里的报错。"""
    with pytest.raises(formats.FormatError) as excinfo:
        formats.read_project("x.txt")
    message = str(excinfo.value)
    assert ".csv" in message and ".xlsx" in message and ".json" in message
    assert isinstance(excinfo.value, ValueError)      # 调用方那套"读不懂=ValueError"照旧管用


def test_the_dialog_filters_come_from_the_registry():
    assert ".csv" in formats.open_filter() and ".xlsx" in formats.open_filter()
    assert formats.save_filter().index(".json") < formats.save_filter().index(".csv")
    assert ".json" not in formats.table_filter(), "导出点表时不该提供工程文件那一项"


# ---------- 三种格式各往返一次 ----------

def test_csv_and_xlsx_round_trip_to_the_same_table(tmp_path):
    """点表格式的判据是**表格形状**（`rows_of`）逐字一致——CSV 与 Excel 都该如此。

    不比整个 `Project`：点表没有 `tick_ms` 那一项（那是工程文件才有的）。
    """
    project = example()
    expected = rows(project)
    for suffix in (".csv", ".xlsx"):
        target = tmp_path / f"t{suffix}"
        formats.write_project(project, target)
        back, _notes = formats.read_project(target)
        assert rows(back) == expected, suffix


def test_json_round_trips_the_whole_project(tmp_path):
    """工程文件是唯一无损的那种：通道类型/串口参数/机架槽位/tick_ms 全在。

    判据用"再写一遍逐字节相同"：`Project`（dataclass）比不出"JSON 里本该没有的东西"，
    而 `json_io.save` 的文本就是它对外的形状（`row` 那种 CSV 行号本来就不进工程文件）。
    """
    project = example()
    first, second = tmp_path / "a.json", tmp_path / "b.json"
    formats.write_project(project, first)
    back, notes = formats.read_project(first)
    formats.write_project(back, second)
    assert first.read_bytes() == second.read_bytes()
    assert notes == []
    assert back.tick_ms == project.tick_ms


# ---------- Excel 那条路的现场形状 ----------

def test_xlsx_finds_the_header_below_a_title_row(tmp_path):
    """表头不必在第一行：现场的表上面常有一行标题；找到之后**行号与 Excel 一致**。"""
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["某某隧道 点表 V1.2"])          # 标题行
    sheet.append([])                              # 空行
    for row in rows(example()):
        sheet.append(list(row))
    sheet.append(["tcp-5020", "D1", "not-an-int", "坏行", "hr", "0", "uint16", "ABCD",
                  "0", "const", "", "false"])
    target = tmp_path / "title.xlsx"
    workbook.save(target)

    with pytest.raises(ValueError) as excinfo:
        formats.read_project(target)
    assert str(excinfo.value).startswith("line 13:"), str(excinfo.value)
    # 第 13 行 = Excel 里那一行（标题 1 + 空行 1 + 表头 1 + 示例 9 行点位 → 坏行落在第 13 行）；
    # 用"先去掉坏行再看读得回来"两条一起钉"行号既对得上、内容也没错位"
    workbook = openpyxl.load_workbook(target)
    workbook.active.delete_rows(13)
    clean = tmp_path / "clean.xlsx"
    workbook.save(clean)
    back, notes = formats.read_project(clean)
    assert rows(back) == rows(example()), "表头前的标题行把内容读错位了"
    assert any("表头在第 3 行" in note for note in notes), notes


def test_xlsx_says_why_a_number_lost_precision(tmp_path):
    """Excel 的数字只有 15 位有效数字：17 位的 uint64 当数字存就丢精度，**要说一句**。

    现场最可能的路径就是"CSV 导进 Excel 改一改、再存成 .xlsx"——那一步 Excel 自己把那串
    数字变成了浮点；不说清楚的话，现场只会看到一句"initial … is not an integer"。
    """
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(list(csv_io.HEADER))
    sheet.append(["tcp-5020", "D1", 11, "累计电量", "hr", 5, "uint64", "ABCD",
                  72623859790382856, "const", "", "false"])
    target = tmp_path / "big.xlsx"
    workbook.save(target)

    project, notes = xlsx_io.read_xlsx_file(target, with_notes=True)
    initial = project.channels[0].devices[0].points[0].initial
    assert initial.startswith("7.262385979038285e+16"), initial
    assert any("15 位有效数字" in note for note in notes), notes


def test_our_own_xlsx_export_writes_such_values_as_text(tmp_path):
    """反过来：我们导出时把那种值写成**文本格**，Excel 不去动它就逐字往返。"""
    import openpyxl

    project = Project(channels=[TcpChannel("tcp-1", port=5020, devices=[
        Device("D1", 11, [Point("累计电量", "hr", 5, "uint64",
                                initial="72623859790382856")])])])
    target = tmp_path / "ours.xlsx"
    xlsx_io.write_xlsx_file(project, target)
    cell = openpyxl.load_workbook(target).active.cell(row=2, column=9)
    assert cell.value == "72623859790382856" and isinstance(cell.value, str)
    back, _notes = formats.read_project(target)
    assert rows(back) == rows(project)



# ---------- 命令行那一路 ----------

def test_the_command_line_entry_loads_an_xlsx_too(tmp_path):
    """命令行（`csv_io_entry.load_point_table`）与界面走**同一张**格式表。

    冻结包那条冒烟（`packaging/smoke.py`）打的就是这条路：源码里不钉住的话，改了界面那一侧
    的过滤器而漏了这里的后缀分派，只有到打包验收才会发现。
    """
    target = tmp_path / "t.xlsx"
    formats.write_project(example(), target)
    project, notes, issues = load_point_table(str(target))
    assert not [issue for issue in issues if issue.level == "error"], \
        [str(issue) for issue in issues]
    assert [(ch.name, ch.port) for ch in project.channels] == [("tcp-5020", 5020)]
    # 演示点表把 host/port 都写在通道名里（`tcp-5020@tcp:0.0.0.0:5020`），所以
    # **没有**"端口按名字尾数猜的"那条 note——两种点表拿到的 notes 都该是空的
    assert notes == [], notes


def test_the_command_line_entry_refuses_an_unknown_suffix(tmp_path):
    """后缀不认识：一行 `[error] [parse]`，不是 traceback（命令行那句人话就是这么来的）。"""
    target = tmp_path / "t.txt"
    target.write_text("channel,device\n", encoding="utf-8")
    project, notes, issues = load_point_table(str(target))
    assert project.channels == [] and notes == []
    assert [issue.level for issue in issues] == ["error"]
    assert ".xlsx" in str(issues[0]), str(issues[0])


# ---------- 图形界面那两条入口 ----------

def test_the_import_entry_reads_xlsx_too(tmp_path):
    """界面的"导入点表"走的还是 `read_point_table`，只是它现在按后缀分派。"""
    target = tmp_path / "t.xlsx"
    formats.write_project(example(), target)
    project, notes = dialogs.read_point_table(str(target))
    assert len(project.channels) == 1 and project.channels[0].port == 5020
    assert rows(project) == rows(example())


def test_save_to_an_xlsx_path_writes_an_excel_file(tmp_path):
    """另存为 `.xlsx` 就是导出一份 Excel 点表（不再是"按 CSV 写出去但叫 xlsx"）。"""
    import openpyxl

    target = tmp_path / "saved.xlsx"
    dialogs.save_project_to(example(), target)
    sheet = openpyxl.load_workbook(target).active
    assert [cell.value for cell in next(sheet.iter_rows())] == csv_io.HEADER
    assert sheet.title == xlsx_io.SHEET_NAME
